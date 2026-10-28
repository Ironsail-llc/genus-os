"""Text out of the document formats an operator actually sends.

One answer for both readers: ``read_file`` (a path the agent was pointed at)
and the Telegram intake (bytes the operator attached). Before this module,
``read_file`` died with a UTF-8 decode error on any binary file and the intake
extracted PDF and nothing else, so a Word document, a spreadsheet or a deck
reached the agent as a path it could not open.

Every parsing library is imported lazily, inside the function that needs it.
An install without one gets :class:`ExtractorUnavailableError` naming the package
to install — never an ImportError at module import, which would take
``read_file`` down for plain text too.
"""

from __future__ import annotations

import io
import mimetypes
import shutil
import subprocess
from pathlib import Path
from typing import Any

__all__ = [
    "XLSX_MAX_ROWS_PER_SHEET",
    "ExtractionError",
    "ExtractorUnavailableError",
    "binary_summary",
    "describe_bytes",
    "document_kind",
    "extract_bytes",
    "extract_path",
    "looks_binary",
]

#: Rows quoted per worksheet. A spreadsheet with a million rows would
#: otherwise be rendered whole only to be cut at the caller's character cap.
XLSX_MAX_ROWS_PER_SHEET = 2_000

#: How long the ``pdftotext`` fallback may run on one file.
_PDFTOTEXT_TIMEOUT_S = 60

_SUFFIX_KIND = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".docm": "docx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".pptx": "pptx",
    ".pptm": "pptx",
}

_MIME_KIND = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
}

_PACKAGE = {"docx": "python-docx", "xlsx": "openpyxl", "pptx": "python-pptx"}


class ExtractionError(Exception):
    """The file could not be turned into text (corrupt, encrypted, ...)."""


class ExtractorUnavailableError(ExtractionError):
    """The library this format needs is not installed."""


def document_kind(suffix: str, mime: str = "") -> str | None:
    """``"pdf"``, ``"docx"``, ``"xlsx"``, ``"pptx"`` or None.

    Suffix first — it is the sender's own label — then MIME, for the
    platforms that report a type and a name without an extension.
    """
    kind = _SUFFIX_KIND.get((suffix or "").lower())
    if kind:
        return kind
    return _MIME_KIND.get((mime or "").strip().lower())


def extract_bytes(raw: bytes, kind: str) -> str:
    """The text of a document of ``kind``. Raises :class:`ExtractionError`."""
    extractor = _EXTRACTORS.get(kind)
    if extractor is None:
        raise ExtractionError(f"no text extractor for {kind!r} files")
    try:
        return extractor(raw)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - every parser has its own error zoo
        raise ExtractionError(f"could not read this {kind} file: {exc}") from exc


def extract_path(path: Path, kind: str) -> str:
    """:func:`extract_bytes` on a file on disk."""
    return extract_bytes(Path(path).read_bytes(), kind)


def describe_bytes(raw: bytes, kind: str) -> str:
    """:func:`extract_bytes` that never raises: a failure becomes a note.

    For the channel layer, where the file is already saved and the worst case
    is an agent told why it got no text.
    """
    try:
        return extract_bytes(raw, kind)
    except ExtractionError as exc:
        return f"[{kind.upper()} text extraction failed: {exc}]"


def looks_binary(sample: bytes) -> bool:
    """A NUL byte in the first block is the cheap, reliable binary tell."""
    return b"\x00" in sample


def binary_summary(path: Path) -> dict[str, Any]:
    """What can honestly be said about a file there is no reader for."""
    path = Path(path)
    mime, _ = mimetypes.guess_type(path.name)
    size = path.stat().st_size
    return {
        "path": str(path),
        "binary": True,
        "size": size,
        "mime": mime or "application/octet-stream",
        "message": (
            f"Binary file ({size} bytes, {mime or 'unknown type'}); there is no text "
            "reader for this format, so its contents were not returned."
        ),
    }


# ── Per-format extractors ────────────────────────────────────────────


def _need(kind: str) -> Any:
    module = {"docx": "docx", "xlsx": "openpyxl", "pptx": "pptx"}[kind]
    try:
        return __import__(module)
    except ImportError as exc:
        raise ExtractorUnavailableError(
            f"reading {kind} files needs the {_PACKAGE[kind]} package "
            f"(pip install {_PACKAGE[kind]})"
        ) from exc


def _pdf(raw: bytes) -> str:
    try:
        import pypdf
    except ImportError:
        return _pdf_via_pdftotext(raw)
    reader = pypdf.PdfReader(io.BytesIO(raw))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Page {i + 1}]\n{text}")
    if pages:
        return "\n\n".join(pages)
    return "[PDF: no extractable text (may be image-based)]"


def _pdf_via_pdftotext(raw: bytes) -> str:
    exe = shutil.which("pdftotext")
    if not exe:
        raise ExtractorUnavailableError(
            "reading pdf files needs the pypdf package (pip install pypdf) "
            "or the pdftotext command (poppler-utils)"
        )
    try:
        done = subprocess.run(  # noqa: S603 - fixed argv, bytes on stdin
            [exe, "-layout", "-", "-"],
            input=raw,
            capture_output=True,
            timeout=_PDFTOTEXT_TIMEOUT_S,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise ExtractionError("pdftotext timed out") from exc
    if done.returncode != 0:
        detail = done.stderr.decode("utf-8", errors="replace").strip()
        raise ExtractionError(f"pdftotext failed: {detail or done.returncode}")
    text = done.stdout.decode("utf-8", errors="replace")
    return text if text.strip() else "[PDF: no extractable text (may be image-based)]"


def _docx(raw: bytes) -> str:
    docx = _need("docx")
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(io.BytesIO(raw))
    parts: list[str] = []
    # Body order, so a table sits between the paragraphs it sat between.
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text
            if text.strip():
                parts.append(text)
        elif tag == "tbl":
            rows = [
                "\t".join(cell.text.strip() for cell in row.cells) for row in Table(child, doc).rows
            ]
            if rows:
                parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _xlsx(raw: bytes) -> str:
    openpyxl = _need("xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for ws in wb.worksheets:
            lines = [f"## Sheet: {ws.title}"]
            kept = 0
            truncated = False
            for row in ws.iter_rows(values_only=True):
                if all(v is None for v in row):
                    continue
                if kept >= XLSX_MAX_ROWS_PER_SHEET:
                    truncated = True
                    break
                lines.append("\t".join("" if v is None else str(v) for v in row))
                kept += 1
            if truncated:
                lines.append(f"[... rows after {XLSX_MAX_ROWS_PER_SHEET} truncated]")
            sheets.append("\n".join(lines))
        return "\n\n".join(sheets)
    finally:
        wb.close()


def _pptx(raw: bytes) -> str:
    pptx = _need("pptx")
    prs = pptx.Presentation(io.BytesIO(raw))
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = [f"[Slide {i}]"]
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False):
                lines.extend(
                    "\t".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows
                )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        slides.append("\n".join(lines))
    return "\n\n".join(slides)


_EXTRACTORS = {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "pptx": _pptx}
