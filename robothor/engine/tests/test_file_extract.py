"""The assistant reads the common file types the operator sends or points it at.

Before this, ``read_file`` called ``Path.read_text()``: a PDF, a Word document
or a spreadsheet died with a UTF-8 decode error, and a Telegram attachment
other than PDF or plain text reached the agent as a path and nothing else.

Fixture files are written programmatically into ``tmp_path`` with the same
libraries the extractor reads them with, so no binary fixture is checked in.
"""

from __future__ import annotations

import builtins
import io
from typing import TYPE_CHECKING
from unittest.mock import MagicMock

import pytest

from robothor.engine import attachments, file_extract
from robothor.engine.tools.handlers import filesystem

if TYPE_CHECKING:
    from pathlib import Path

# ── Fixture builders ─────────────────────────────────────────────────


def make_docx(path: Path) -> Path:
    docx = pytest.importorskip("docx")
    doc = docx.Document()
    doc.add_paragraph("Quarterly plan for Alice")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Budget"
    table.cell(1, 0).text = "Bob"
    table.cell(1, 1).text = "1200"
    doc.add_paragraph("Closing paragraph")
    doc.save(str(path))
    return path


def make_xlsx(path: Path, *, rows: int = 2) -> Path:
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Region", "Amount"])
    for i in range(rows - 1):
        ws.append([f"R{i}", 100 + i])
    second = wb.create_sheet("Notes")
    second.append(["remember the milk"])
    wb.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation()
    for title, body in (("Kickoff", "Agenda item one"), ("Roadmap", "Ship the reader")):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(str(path))
    return path


def make_pdf_bytes(text: str = "Hello PDF world") -> bytes:
    """A minimal one-page PDF with a real text stream, built by hand."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def hide_module(monkeypatch, name: str) -> None:
    """Make ``import name`` raise ImportError, as on an install without it."""
    real_import = builtins.__import__

    def fake_import(mod, *args, **kwargs):
        if mod == name or mod.startswith(name + "."):
            raise ImportError(f"No module named {name!r}")
        return real_import(mod, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)


async def read_file(path, workspace):
    ctx = MagicMock()
    ctx.workspace = str(workspace)
    return await filesystem.HANDLERS["read_file"]({"path": str(path)}, ctx)


# ── The shared extractor ─────────────────────────────────────────────


class TestDocumentKind:
    @pytest.mark.parametrize(
        ("suffix", "mime", "expected"),
        [
            (".pdf", "", "pdf"),
            ("", "application/pdf", "pdf"),
            (".docx", "", "docx"),
            (".DOCX", "", "docx"),
            (".xlsx", "", "xlsx"),
            (".xlsm", "", "xlsx"),
            (".pptx", "", "pptx"),
            (
                "",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "docx",
            ),
            ("", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
            (
                "",
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "pptx",
            ),
            (".txt", "text/plain", None),
            (".zip", "application/zip", None),
        ],
    )
    def test_kind(self, suffix, mime, expected) -> None:
        assert file_extract.document_kind(suffix, mime) == expected


class TestExtract:
    def test_docx_paragraphs_and_tables_in_order(self, tmp_path) -> None:
        text = file_extract.extract_bytes(make_docx(tmp_path / "a.docx").read_bytes(), "docx")
        assert "Quarterly plan for Alice" in text
        assert "Owner\tBudget" in text
        assert "Bob\t1200" in text
        assert text.index("Quarterly plan") < text.index("Owner") < text.index("Closing")

    def test_xlsx_every_sheet_tab_separated(self, tmp_path) -> None:
        text = file_extract.extract_bytes(make_xlsx(tmp_path / "a.xlsx").read_bytes(), "xlsx")
        assert "Revenue" in text
        assert "Region\tAmount" in text
        assert "R0\t100" in text
        assert "Notes" in text
        assert "remember the milk" in text

    def test_xlsx_rows_are_capped(self, tmp_path, monkeypatch) -> None:
        monkeypatch.setattr(file_extract, "XLSX_MAX_ROWS_PER_SHEET", 5)
        raw = make_xlsx(tmp_path / "big.xlsx", rows=50).read_bytes()
        text = file_extract.extract_bytes(raw, "xlsx")
        assert "R3\t103" in text
        assert "R40" not in text
        assert "truncated" in text.lower()

    def test_pptx_text_per_slide(self, tmp_path) -> None:
        text = file_extract.extract_bytes(make_pptx(tmp_path / "a.pptx").read_bytes(), "pptx")
        assert "Slide 1" in text and "Slide 2" in text
        assert "Kickoff" in text and "Agenda item one" in text
        assert text.index("Kickoff") < text.index("Roadmap")

    def test_pdf_with_pypdf(self) -> None:
        pytest.importorskip("pypdf")
        text = file_extract.extract_bytes(make_pdf_bytes(), "pdf")
        assert "Hello PDF world" in text

    def test_pdf_falls_back_to_pdftotext(self, monkeypatch) -> None:
        import shutil

        if not shutil.which("pdftotext"):
            pytest.skip("pdftotext not installed")
        hide_module(monkeypatch, "pypdf")
        text = file_extract.extract_bytes(make_pdf_bytes(), "pdf")
        assert "Hello PDF world" in text

    def test_pdf_with_nothing_installed_says_so(self, monkeypatch) -> None:
        hide_module(monkeypatch, "pypdf")
        monkeypatch.setattr(file_extract.shutil, "which", lambda _name: None)
        with pytest.raises(file_extract.ExtractorUnavailableError, match="pypdf"):
            file_extract.extract_bytes(make_pdf_bytes(), "pdf")

    @pytest.mark.parametrize(
        ("kind", "module", "package"),
        [
            ("docx", "docx", "python-docx"),
            ("xlsx", "openpyxl", "openpyxl"),
            ("pptx", "pptx", "python-pptx"),
        ],
    )
    def test_missing_library_is_a_clear_message(self, monkeypatch, kind, module, package) -> None:
        hide_module(monkeypatch, module)
        with pytest.raises(file_extract.ExtractorUnavailableError, match=package):
            file_extract.extract_bytes(b"PK\x03\x04", kind)

    def test_corrupt_file_is_an_extraction_error_not_a_crash(self) -> None:
        with pytest.raises(file_extract.ExtractionError):
            file_extract.extract_bytes(b"definitely not a zip", "docx")

    def test_describe_extracts_or_explains(self) -> None:
        # The channel layer never raises: a failure becomes a bracketed note.
        out = file_extract.describe_bytes(b"definitely not a zip", "docx")
        assert out.startswith("[") and "docx" in out.lower()


class TestBinarySummary:
    def test_summary_has_size_and_mime(self, tmp_path) -> None:
        p = tmp_path / "bundle.zip"
        p.write_bytes(b"PK\x03\x04" + b"\x00" * 100)
        out = file_extract.binary_summary(p)
        assert out["size"] == 104
        assert out["mime"] == "application/zip"
        assert out["binary"] is True


# ── read_file ────────────────────────────────────────────────────────


class TestReadFile:
    @pytest.mark.asyncio
    async def test_docx(self, tmp_path) -> None:
        out = await read_file(make_docx(tmp_path / "plan.docx"), tmp_path)
        assert "error" not in out, out
        assert "Quarterly plan for Alice" in out["content"]
        assert out["format"] == "docx"
        assert out["chars"] == len(out["content"])

    @pytest.mark.asyncio
    async def test_xlsx(self, tmp_path) -> None:
        out = await read_file(make_xlsx(tmp_path / "rev.xlsx"), tmp_path)
        assert "Region\tAmount" in out["content"]

    @pytest.mark.asyncio
    async def test_pptx(self, tmp_path) -> None:
        out = await read_file(make_pptx(tmp_path / "deck.pptx"), tmp_path)
        assert "Roadmap" in out["content"]

    @pytest.mark.asyncio
    async def test_pdf(self, tmp_path) -> None:
        pytest.importorskip("pypdf")
        p = tmp_path / "doc.pdf"
        p.write_bytes(make_pdf_bytes())
        out = await read_file(p, tmp_path)
        assert "Hello PDF world" in out["content"]

    @pytest.mark.asyncio
    async def test_extracted_text_uses_the_same_cap(self, tmp_path, monkeypatch) -> None:
        monkeypatch.setattr(
            file_extract, "extract_path", lambda _p, _k, max_chars=None: "y" * 60_000
        )
        p = tmp_path / "huge.docx"
        p.write_bytes(b"PK")
        out = await read_file(p, tmp_path)
        assert len(out["content"]) == 50_000
        assert out["chars"] == 60_000

    @pytest.mark.asyncio
    async def test_an_image_points_to_view_image(self, tmp_path) -> None:
        p = tmp_path / "photo.png"
        p.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
        out = await read_file(p, tmp_path)
        assert "view_image" in str(out)
        assert "decode" not in str(out).lower()

    @pytest.mark.asyncio
    async def test_other_binary_is_a_summary_not_a_decode_error(self, tmp_path) -> None:
        p = tmp_path / "bundle.bin"
        p.write_bytes(b"\x00\xff\xfe\x01" * 64)
        out = await read_file(p, tmp_path)
        assert "error" not in out, out
        assert out["binary"] is True
        assert out["size"] == 256
        assert "decode" not in str(out).lower()

    @pytest.mark.asyncio
    async def test_plain_text_is_unchanged(self, tmp_path) -> None:
        p = tmp_path / "notes.txt"
        p.write_text("hello there")
        out = await read_file(p, tmp_path)
        assert out == {"content": "hello there", "path": str(p), "chars": 11}

    @pytest.mark.asyncio
    async def test_a_corrupt_docx_is_an_error_message(self, tmp_path) -> None:
        p = tmp_path / "broken.docx"
        p.write_bytes(b"not a zip at all")
        out = await read_file(p, tmp_path)
        assert "error" in out
        assert "docx" in out["error"].lower()

    @pytest.mark.asyncio
    async def test_missing_file_is_still_an_error(self, tmp_path) -> None:
        out = await read_file(tmp_path / "nope.docx", tmp_path)
        assert "error" in out

    @pytest.mark.asyncio
    async def test_secret_path_is_still_refused(self, tmp_path) -> None:
        # A private key: refused outright, whatever happens to dotenv files.
        p = tmp_path / "id_rsa"
        p.write_text("-----BEGIN OPENSSH PRIVATE KEY-----\nsuper-secret\n")
        out = await read_file(p, tmp_path)
        assert "error" in out
        assert "super-secret" not in str(out)
        assert "PRIVATE KEY" not in str(out)


# ── Attachments ──────────────────────────────────────────────────────


class TestExtractable:
    @pytest.mark.parametrize("suffix", [".docx", ".xlsx", ".pptx"])
    def test_office_formats_extract(self, suffix) -> None:
        assert attachments.extractable(suffix) == suffix[1:]

    def test_office_mime_alone_extracts(self) -> None:
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert attachments.extractable("", mime) == "docx"

    def test_pdf_and_text_unchanged(self) -> None:
        assert attachments.extractable(".pdf") == "pdf"
        assert attachments.extractable(".md") == "text"
        assert attachments.extractable(".zip") is None


class TestPdfHelper:
    @pytest.mark.asyncio
    async def test_the_telegram_pdf_helper_uses_the_shared_extractor(self) -> None:
        pytest.importorskip("pypdf")
        from robothor.engine.telegram_attachments import _extract_pdf_text

        assert "Hello PDF world" in await _extract_pdf_text(make_pdf_bytes())


# ── Review follow-ups: budget, .docm, pdftotext, nested content ──────


class FakePage:
    def __init__(self, text: str, visited: list[int], index: int) -> None:
        self._text = text
        self._visited = visited
        self._index = index

    def extract_text(self) -> str:
        self._visited.append(self._index)
        return self._text


def fake_pdf_reader(monkeypatch, pages: list[str]) -> list[int]:
    """Swap pypdf.PdfReader for one that records which pages were read."""
    pypdf = pytest.importorskip("pypdf")
    visited: list[int] = []

    class Reader:
        def __init__(self, _stream) -> None:
            self.pages = [FakePage(t, visited, i) for i, t in enumerate(pages)]

    monkeypatch.setattr(pypdf, "PdfReader", Reader)
    return visited


def make_docm_bytes(path: Path) -> bytes:
    """A real macro-enabled Word package: a docx with the macro main type."""
    import zipfile

    make_docx(path)
    out = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument."
                    b"wordprocessingml.document.main+xml",
                    b"application/vnd.ms-word.document.macroEnabled.main+xml",
                )
            dst.writestr(item, data)
    return out.getvalue()


class TestBudget:
    def test_pdf_stops_reading_pages_once_the_budget_is_reached(self, monkeypatch) -> None:
        visited = fake_pdf_reader(monkeypatch, ["x" * 1_000] * 1_000)
        text = file_extract.extract_bytes(b"%PDF", "pdf", max_chars=5_000)
        assert len(visited) < 10
        assert text.startswith("[Page 1]")
        assert file_extract.STOPPED_EARLY in text

    def test_pdf_without_a_budget_reads_everything(self, monkeypatch) -> None:
        visited = fake_pdf_reader(monkeypatch, ["x" * 100] * 50)
        text = file_extract.extract_bytes(b"%PDF", "pdf")
        assert len(visited) == 50
        assert file_extract.STOPPED_EARLY not in text

    def test_a_document_that_fits_is_not_marked(self, monkeypatch) -> None:
        fake_pdf_reader(monkeypatch, ["short"] * 3)
        text = file_extract.extract_bytes(b"%PDF", "pdf", max_chars=5_000)
        assert file_extract.STOPPED_EARLY not in text

    def test_budgeted_text_is_cut_at_the_budget(self, monkeypatch) -> None:
        fake_pdf_reader(monkeypatch, ["x" * 1_000] * 100)
        text = file_extract.extract_bytes(b"%PDF", "pdf", max_chars=2_500)
        body, _, note = text.partition("\n\n" + file_extract.STOPPED_EARLY)
        assert len(body) == 2_500
        assert note

    def test_docx_stops_at_the_budget(self, tmp_path) -> None:
        docx = pytest.importorskip("docx")
        doc = docx.Document()
        for i in range(400):
            doc.add_paragraph(f"paragraph {i} " + "w" * 100)
        doc.save(str(tmp_path / "long.docx"))
        raw = (tmp_path / "long.docx").read_bytes()
        text = file_extract.extract_bytes(raw, "docx", max_chars=1_000)
        assert "paragraph 0 " in text
        assert "paragraph 300 " not in text
        assert file_extract.STOPPED_EARLY in text

    def test_pptx_stops_at_the_budget(self, tmp_path) -> None:
        pptx = pytest.importorskip("pptx")
        prs = pptx.Presentation()
        for i in range(60):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = f"Title {i}"
            slide.placeholders[1].text = "b" * 200
        prs.save(str(tmp_path / "long.pptx"))
        raw = (tmp_path / "long.pptx").read_bytes()
        text = file_extract.extract_bytes(raw, "pptx", max_chars=1_000)
        assert "Title 0" in text
        assert "Title 50" not in text
        assert file_extract.STOPPED_EARLY in text

    def test_xlsx_stops_before_later_sheets(self, tmp_path) -> None:
        openpyxl = pytest.importorskip("openpyxl")
        wb = openpyxl.Workbook()
        wb.active.title = "First"
        for i in range(500):
            wb.active.append([f"row{i}", "v" * 50])
        wb.create_sheet("Last").append(["never reached"])
        wb.save(str(tmp_path / "wide.xlsx"))
        raw = (tmp_path / "wide.xlsx").read_bytes()
        text = file_extract.extract_bytes(raw, "xlsx", max_chars=2_000)
        assert "row0" in text
        assert "row400" not in text
        assert "never reached" not in text
        assert file_extract.STOPPED_EARLY in text

    @pytest.mark.asyncio
    async def test_read_file_passes_its_cap(self, tmp_path, monkeypatch) -> None:
        seen = {}

        def fake(_p, _k, *, max_chars=None):
            seen["max_chars"] = max_chars
            return "text"

        monkeypatch.setattr(file_extract, "extract_path", fake)
        p = tmp_path / "a.docx"
        p.write_bytes(b"PK")
        await read_file(p, tmp_path)
        assert seen["max_chars"] == 50_000

    @pytest.mark.asyncio
    async def test_read_file_says_when_it_stopped_early(self, tmp_path, monkeypatch) -> None:
        fake_pdf_reader(monkeypatch, ["x" * 1_000] * 200)
        p = tmp_path / "big.pdf"
        p.write_bytes(b"%PDF")
        out = await read_file(p, tmp_path)
        assert out["truncated"] is True
        assert file_extract.STOPPED_EARLY in out["content"]


class TestDocm:
    def test_python_docx_alone_refuses_a_docm(self, tmp_path) -> None:
        docx = pytest.importorskip("docx")
        with pytest.raises(ValueError, match="not a Word file"):
            docx.Document(io.BytesIO(make_docm_bytes(tmp_path / "m.docx")))

    def test_docm_is_read(self, tmp_path) -> None:
        raw = make_docm_bytes(tmp_path / "m.docx")
        assert file_extract.document_kind(".docm") == "docx"
        text = file_extract.extract_bytes(raw, "docx")
        assert "Quarterly plan for Alice" in text
        assert "Bob\t1200" in text

    def test_macro_office_mime_types_route(self) -> None:
        assert (
            file_extract.document_kind("", "application/vnd.ms-word.document.macroEnabled.12")
            == "docx"
        )
        assert (
            file_extract.document_kind(
                "", "application/vnd.ms-powerpoint.presentation.macroEnabled.12"
            )
            == "pptx"
        )


class TestPdfFallback:
    def _pdftotext(self, monkeypatch, stdout: bytes = b"from poppler", rc: int = 0) -> list:
        calls: list = []
        monkeypatch.setattr(file_extract.shutil, "which", lambda _n: "/usr/bin/pdftotext")

        def run(argv, **_kw):
            calls.append(argv)
            return MagicMock(returncode=rc, stdout=stdout, stderr=b"boom")

        monkeypatch.setattr(file_extract.subprocess, "run", run)
        return calls

    def _pypdf_raises(self, monkeypatch) -> None:
        pypdf = pytest.importorskip("pypdf")

        def broken(_stream):
            raise ValueError("malformed xref")

        monkeypatch.setattr(pypdf, "PdfReader", broken)

    def test_falls_back_when_pypdf_raises(self, monkeypatch) -> None:
        self._pypdf_raises(monkeypatch)
        calls = self._pdftotext(monkeypatch)
        assert file_extract.extract_bytes(b"%PDF", "pdf") == "from poppler"
        assert calls

    def test_falls_back_when_pypdf_finds_no_text(self, monkeypatch) -> None:
        fake_pdf_reader(monkeypatch, ["", "  "])
        calls = self._pdftotext(monkeypatch)
        assert file_extract.extract_bytes(b"%PDF", "pdf") == "from poppler"
        assert calls

    def test_no_text_from_either_is_the_image_note(self, monkeypatch) -> None:
        fake_pdf_reader(monkeypatch, [""])
        self._pdftotext(monkeypatch, stdout=b"   ")
        assert "no extractable text" in file_extract.extract_bytes(b"%PDF", "pdf")

    def test_pypdf_raises_and_no_binary_is_pypdfs_error(self, monkeypatch) -> None:
        self._pypdf_raises(monkeypatch)
        monkeypatch.setattr(file_extract.shutil, "which", lambda _n: None)
        with pytest.raises(file_extract.ExtractionError, match="malformed xref"):
            file_extract.extract_bytes(b"%PDF", "pdf")

    def test_no_text_and_no_binary_is_the_image_note(self, monkeypatch) -> None:
        fake_pdf_reader(monkeypatch, [""])
        monkeypatch.setattr(file_extract.shutil, "which", lambda _n: None)
        assert "no extractable text" in file_extract.extract_bytes(b"%PDF", "pdf")

    def test_both_failing_names_both(self, monkeypatch) -> None:
        self._pypdf_raises(monkeypatch)
        self._pdftotext(monkeypatch, rc=1)
        with pytest.raises(file_extract.ExtractionError) as err:
            file_extract.extract_bytes(b"%PDF", "pdf")
        assert "malformed xref" in str(err.value)
        assert "pdftotext" in str(err.value)

    def test_pdftotext_output_respects_the_budget(self, monkeypatch) -> None:
        self._pypdf_raises(monkeypatch)
        self._pdftotext(monkeypatch, stdout=b"z" * 10_000)
        text = file_extract.extract_bytes(b"%PDF", "pdf", max_chars=1_000)
        assert text.startswith("z" * 1_000)
        assert file_extract.STOPPED_EARLY in text


class TestNestedContent:
    def test_pptx_grouped_shapes_are_read(self, tmp_path) -> None:
        pptx = pytest.importorskip("pptx")
        from pptx.util import Inches

        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        outer = slide.shapes.add_group_shape()
        box = outer.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        box.text_frame.text = "Grouped callout"
        inner = outer.shapes.add_group_shape()
        deep = inner.shapes.add_textbox(Inches(3), Inches(1), Inches(2), Inches(1))
        deep.text_frame.text = "Nested deeper"
        prs.save(str(tmp_path / "g.pptx"))
        text = file_extract.extract_bytes((tmp_path / "g.pptx").read_bytes(), "pptx")
        assert "Grouped callout" in text
        assert "Nested deeper" in text

    def test_docx_body_content_controls_are_read(self, tmp_path) -> None:
        docx = pytest.importorskip("docx")
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        doc = docx.Document()
        doc.add_paragraph("Before the control")
        sdt = parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtPr/><w:sdtContent>"
            "<w:p><w:r><w:t>Cover page title</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt>"
        )
        doc.element.body.insert(1, sdt)
        doc.add_paragraph("After the control")
        doc.save(str(tmp_path / "sdt.docx"))
        text = file_extract.extract_bytes((tmp_path / "sdt.docx").read_bytes(), "docx")
        assert "Cover page title" in text
        assert text.index("Before") < text.index("Cover page") < text.index("After")
